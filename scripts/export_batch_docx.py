# -*- coding: utf-8 -*-
"""
Script tự động kết xuất trọn bộ 5 File Word (.docx) chuyên dụng cho sản xuất hàng loạt chuẩn V10
Skill: sport-girl-cycling-vlog
Quy chuẩn: 100% văn bản thuần túy (Plain Text Paragraphs), KHÔNG sử dụng Table ở bất kỳ file nào.
Chuẩn Tool Batch: File 01, 02, 03 mỗi cảnh là đúng 1 đoạn văn liên tục không nhãn thừa.
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_styled_title(doc, title_text, subtitle_text, note_text=None):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run(title_text)
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(14, 165, 233)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    if note_text:
        p_sub.paragraph_format.space_after = Pt(4)
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_note = p_note.add_run(note_text)
        run_note.font.name = "Segoe UI"
        run_note.font.size = Pt(9.5)
        run_note.font.color.rgb = RGBColor(245, 158, 11)
        p_note.paragraph_format.space_after = Pt(16)
    else:
        p_sub.paragraph_format.space_after = Pt(14)

def add_clean_paragraph(doc, text_content, font_name="Consolas", font_size=10, is_bold=False):
    """Thêm đúng 1 đoạn văn liên tục không nhãn thừa cho tool batch"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text_content)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = is_bold
    run.font.color.rgb = RGBColor(30, 41, 59)

def export_all_docx(project_dir, project_data):
    """
    Kết xuất trọn bộ 5 File Word (.docx) chuẩn V10 Tool Batch
    """
    export_dir = os.path.join(project_dir, "export_batch")
    os.makedirs(export_dir, exist_ok=True)

    # File 00: Hồ sơ tham chiếu gốc
    doc0 = Document()
    add_styled_title(doc0, "HỒ SƠ THAM CHIẾU GỐC (MASTER REFERENCE ASSETS)",
                     f"Dự án: {project_data.get('title', 'Sport Girl Cycling Vlog')}",
                     "⚠️ BƯỚC 0 BẮT BUỘC: Nạp ảnh Character Turnaround làm Image/Character Reference cho Tool AI.")
    if 'char_sheet' in project_data:
        p_h = doc0.add_heading(level=2)
        p_h.add_run("1. Master Character Model Sheet")
        add_clean_paragraph(doc0, project_data['char_sheet'])
    if 'loc_sheet' in project_data:
        p_h = doc0.add_heading(level=2)
        p_h.add_run("2. Master Location Sheet")
        add_clean_paragraph(doc0, project_data['loc_sheet'])
    doc0.save(os.path.join(export_dir, "00_HO_SO_THAM_CHIEU_GOC.docx"))

    # File 01: Prompts tạo ảnh hàng loạt (1 đoạn văn/cảnh)
    doc1 = Document()
    add_styled_title(doc1, "PROMPTS TẠO ẢNH HÀNG LOẠT (BATCH IMAGE PROMPTS)",
                     f"Dự án: {project_data.get('title', 'Sport Girl Cycling Vlog')} • Chuẩn 9:16 Real Human DSLR",
                     "Định dạng Tool Batch: 100% văn bản thuần túy không nhãn thừa. Mỗi cảnh là đúng 1 đoạn văn liên tục.")
    for p in project_data.get('image_prompts', []):
        add_clean_paragraph(doc1, p)
    doc1.save(os.path.join(export_dir, "01_PROMPTS_ANH_HANG_LOAT.docx"))

    # File 02: Prompts tạo video hàng loạt (1 đoạn văn/cảnh)
    doc2 = Document()
    add_styled_title(doc2, "PROMPTS TẠO VIDEO HÀNG LOẠT (BATCH MEGA PROMPTS)",
                     f"Dự án: {project_data.get('title', 'Sport Girl Cycling Vlog')} • Chuẩn Tagged Blocks Bắt Buộc",
                     "Định dạng Tool Batch: 100% văn bản thuần túy không nhãn thừa. Mỗi cảnh là đúng 1 đoạn văn liên tục không ngắt dòng bên trong.")
    for vp in project_data.get('video_prompts', []):
        add_clean_paragraph(doc2, vp)
    doc2.save(os.path.join(export_dir, "02_PROMPTS_VIDEO_HANG_LOAT.docx"))

    # File 03: Prompts voice hàng loạt (1 đoạn văn/cảnh, thuần thoại)
    doc3 = Document()
    add_styled_title(doc3, "KỊCH BẢN LỜI THOẠI & ÂM THANH HÀNG LOẠT (BATCH TTS SCRIPTS)",
                     f"Dự án: {project_data.get('title', 'Sport Girl Cycling Vlog')}",
                     "Định dạng Tool Batch: Lời thoại sạch 100%, không nhãn người nói, không chú thích SFX.")
    for vl in project_data.get('voice_prompts', []):
        add_clean_paragraph(doc3, vl, font_name="Segoe UI", font_size=11.5)
    doc3.save(os.path.join(export_dir, "03_PROMPTS_VOICE_HANG_LOAT.docx"))

    # File 04: Metadata SEO
    doc4 = Document()
    add_styled_title(doc4, "GÓI TỐI ƯU SEO & ĐĂNG TẢI VIRAL ĐA KÊNH",
                     f"Dự án: {project_data.get('title', 'Sport Girl Cycling Vlog')} • Chuẩn Đăng Tải Shorts / TikTok / Reels")
    seo = project_data.get('seo', {})
    if 'thumbnail_prompt' in seo:
        p_h = doc4.add_heading(level=2)
        p_h.add_run("1. Prompt Tạo Ảnh Bìa Thumbnail (9:16)")
        add_clean_paragraph(doc4, seo['thumbnail_prompt'])
    if 'overlay_text' in seo:
        p_h = doc4.add_heading(level=2)
        p_h.add_run("2. Chữ Dán Thumbnail (Overlay Text)")
        add_clean_paragraph(doc4, seo['overlay_text'], font_name="Segoe UI", font_size=11)
    if 'headlines' in seo:
        p_h = doc4.add_heading(level=2)
        p_h.add_run("3. Tiêu Đề Đa Nền Tảng (A/B Testing Headlines)")
        add_clean_paragraph(doc4, seo['headlines'], font_name="Segoe UI", font_size=11)
    if 'caption' in seo:
        p_h = doc4.add_heading(level=2)
        p_h.add_run("4. Caption & Hashtag 3 Tầng")
        add_clean_paragraph(doc4, seo['caption'], font_name="Segoe UI", font_size=10.5)
    doc4.save(os.path.join(export_dir, "04_METADATA_SEO.docx"))

    print(f"[OK] Đã xuất thành công trọn bộ 5 file .docx chuẩn Tool Batch tại: {export_dir}")
    return export_dir

if __name__ == "__main__":
    print("Script export_batch_docx.py cho skill sport-girl-cycling-vlog đã sẵn sàng theo chuẩn V10.")
